"""Compile a bounded, whitelisted text format to the existing DOCX renderer."""
from __future__ import annotations

import os
import re
import uuid
import shutil
from typing import Any

from . import report_template as rt
from .report_document import REPORT_FIELDS, FINDING_FIELDS

REPORT_TAGS = tuple('report.' + key for key in REPORT_FIELDS)
STAT_TAGS = tuple('stat.' + key for key in ('total', 'critical', 'high', 'medium', 'low', 'info'))
FINDING_TAGS = tuple('finding.' + key for key in FINDING_FIELDS)
TOKENS = REPORT_TAGS + STAT_TAGS + FINDING_TAGS + ('finding.screenshots',)
START, END = '{{#findings}}', '{{/findings}}'


def compile_document(body: Any):
    """No user-authored Jinja is ever evaluated. Only compiler-generated tags exist."""
    from docx import Document
    from docx.shared import Mm, Pt
    if not isinstance(body, str) or not body.strip() or len(body) > 100000:
        raise ValueError('模板正文不能为空且不得超过100000字符')
    if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', body):
        raise ValueError('模板包含不支持的控制符或模板语句，请使用预设标签')
    lines = body.replace('\r\n', '\n').replace('\r', '\n').split('\n')
    if len(lines) > 2000 or body.count('{{') > 1000:
        raise ValueError('模板行数或标签数量过多')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(25)
    doc.styles['Normal'].font.size = Pt(11)
    schema = {'version': 1, 'scalars': [], 'loops': [], 'images': [], 'fixed_kept': []}
    inside = used = False
    for line in lines:
        trimmed = line.strip()
        if trimmed == START:
            if inside or used:
                raise ValueError('仅支持一个漏洞循环，不能嵌套或重复')
            inside = used = True
            doc.add_paragraph('{%p for item in findings %}')
            continue
        if trimmed == END:
            if not inside:
                raise ValueError('漏洞循环结束标签没有对应起始标签')
            inside = False
            doc.add_paragraph('{%p endfor %}')
            continue
        if trimmed == '{{finding.screenshots}}':
            if not inside:
                raise ValueError('漏洞图片标签须放在漏洞循环内')
            doc.add_paragraph('{%p for shot in item.shots %}')
            doc.add_paragraph('{{ shot }}')
            doc.add_paragraph('{%p endfor %}')
            continue
        if any(tag in line for tag in ('{%', '{#', '%}', '#}')):
            raise ValueError('不支持自定义模板语句，请使用预设标签')
        def replace(match):
            token = match.group(1).strip()
            if token not in TOKENS or token == 'finding.screenshots':
                raise ValueError('不支持的标签：{{' + token[:80] + '}}；图片和循环标签须单独占一行')
            if token.startswith('finding.'):
                if not inside:
                    raise ValueError('漏洞字段须放在 {{#findings}} 与 {{/findings}} 之间')
                return '{{ item.' + token.split('.', 1)[1] + ' }}'
            placeholder = 'manual_' + token.replace('.', '_')
            if not any(s['placeholder'] == placeholder for s in schema['scalars']):
                schema['scalars'].append({'placeholder': placeholder, 'semantic': token})
            return '{{ ' + placeholder + ' }}'
        # Validate all delimiters before replacing them with our own tags.
        rest = re.sub(r'\{\{([^{}]+)\}\}', '', line)
        if '{{' in rest or '}}' in rest:
            raise ValueError('标签括号不完整，请使用 {{字段名}}')
        text = re.sub(r'\{\{([^{}]+)\}\}', replace, line)
        heading = re.match(r'^(#{1,3})\s+(.+)$', text)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)))
        else:
            doc.add_paragraph(text)
    if inside:
        raise ValueError('缺少漏洞循环结束标签 {{/findings}}')
    return doc, schema


def save(payload: Any, created_by: str = '') -> dict:
    """New immutable file + revision CAS: no process cache or in-place DOCX writes."""
    from bson import ObjectId
    if not isinstance(payload, dict):
        return {'error': '模板数据必须为对象'}
    if not payload.get('template_id'):
        return {'error':'新建报告模板请先上传基础Word文档，再选择人工标注或AI学习'}
    name, body = payload.get('name'), payload.get('body')
    if not isinstance(name, str) or not name.strip() or len(name.strip()) > 120:
        return {'error': '模板名称须为1至120字符'}
    try:
        document, schema = compile_document(body)
    except (ValueError, ImportError) as exc:
        return {'error': str(exc)}
    tid = payload.get('template_id', '')
    if not isinstance(tid, str) or (tid and not ObjectId.is_valid(tid)):
        return {'error': '无效的人工模板标识'}
    coll = rt.repo_coll()
    current = coll.find_one({'_id': ObjectId(tid)}) if tid else None
    if tid and not current:
        return {'error': '模板不存在'}
    if current and (current.get('builtin') or current.get('editor_mode') != 'manual'):
        return {'error': '仅可编辑人工模板；内置和AI学习模板请保留或另建人工模板'}
    revision = payload.get('revision')
    if current and (isinstance(revision, bool) or not isinstance(revision, int) or revision != current.get('revision')):
        return {'error': '模板已被其他操作更新，请重新打开后编辑'}
    oid = current['_id'] if current else ObjectId()
    tid = str(oid)
    root = rt._tpl_root(tid)
    os.makedirs(root, exist_ok=True)
    path = os.path.join(root, 'manual_' + uuid.uuid4().hex + '.docx')
    try:
        document.save(path)
        origin_path=current.get('origin_path') if current else None
        if current and not origin_path:
            origin_file=os.path.join(root,'legacy_origin_'+uuid.uuid4().hex+'.docx')
            shutil.copyfile(rt._abs_template_path(current['template_path']),origin_file)
            origin_path=os.path.relpath(origin_file,rt.template_dir()).replace(chr(92),'/')
        now = rt._now()
        record = {'name': name.strip(), 'status': 'review', 'need_review':True, 'editor_mode': 'manual',
                  'manual_body': body, 'revision': (revision + 1 if current else 1),
                  'schema': schema, 'template_path': os.path.relpath(path, rt.template_dir()).replace('\\', '/'),
                  'source_filename': 'manual.docx', 'learn_tokens': 0, 'update_date': now}
        if origin_path:record['origin_path']=origin_path
        if current:
            result = coll.update_one({'_id': oid, 'editor_mode': 'manual', 'builtin': {'$ne': True}, 'revision': revision}, {'$set': record})
            if not result.matched_count:
                os.remove(path)
                return {'error': '模板已被其他操作更新或删除，请重新加载'}
        else:
            coll.insert_one(dict(record, _id=oid, created_by=created_by, save_date=now))
        return {'ok': True, 'template_id': tid, 'status': 'review', 'revision': record['revision']}
    except Exception as exc:
        # Keep an immutable file if a database response is lost after a possible commit.
        rt.logger.warning('manual template save failed: %s', type(exc).__name__)
        return {'error': '模板保存失败，请刷新列表确认后重试'}
