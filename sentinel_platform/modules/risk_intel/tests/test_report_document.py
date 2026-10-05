"""NCC exports: isolated repository/files, no real report data or network."""
import copy
import io
import os
import tempfile
import unittest
import zipfile
from unittest.mock import patch

from docx import Document
from PIL import Image
from sentinel_platform.core import set_repo
from sentinel_platform.modules.risk_intel import report_template as rt
from sentinel_platform.modules.risk_intel import report_document as rd
from sentinel_platform.modules.risk_intel import report_template_ncc as ncc
from sentinel_platform.modules.risk_intel import asset_intel
from sentinel_platform.modules.risk_intel.tests.test_report_template import _Repo


def png(size=(1400, 500), fmt="PNG"):
    buf = io.BytesIO()
    Image.new("RGB", size, (80, 120, 160)).save(buf, format=fmt)
    return buf.getvalue()


class ReportDocumentTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.repo = _Repo()
        set_repo(self.repo)
        self.addCleanup(set_repo, None)
        self.root = self.temp.name
        for module, attr in ((rt, "template_dir"), (rt, "image_dir")):
            p = patch.object(module, attr, return_value=self.root)
            p.start()
            self.addCleanup(p.stop)
        self.repo.collection("intel_finding").insert_one({
            "_id": "fixture", "severity": "high", "vuln_type": "权限校验缺失", "unit": "示例单位",
            "target": "https://example.invalid/", "verified": True,
            "poc": "GET /example?x=<value>&a=1 HTTP/1.1", "evidence": "E" * 2501 + "END",
            "verify_method": "在测试环境确认响应", "save_date": "2026-09-25 10:00:00",
        })
        self.assertEqual(ncc.seed(), 1)

    def generate(self):
        result = rt.generate_from_template(ncc.TEMPLATE_ID, "finding", "fixture")
        self.assertTrue(result.get("ok"), result)
        return self.repo.collection("pentest_report").find_one({"_id": result["report_id"]})

    def test_image_upload_normalizes_webp_and_rejects_invalid(self):
        result = rt.save_finding_shot("fixture", png(fmt="WEBP"), "webp")
        self.assertTrue(result.get("ok"), result)
        self.assertTrue(result["name"].endswith(".png"))
        self.assertIn("error", rt.save_finding_shot("fixture", b"not an image"))
        self.assertIn("error", rt.save_finding_shot("../escape", png()))

    def test_ncc_preserves_text_and_page_geometry(self):
        report = self.generate()
        doc = Document(os.path.join(self.root, report["docx_path"]))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("<value>&a=1", text)
        self.assertIn("E" * 2501 + "END", text)
        self.assertIn("5.修复建议", text)
        self.assertNotIn("{{", text)
        self.assertNotIn("[例如", text)
        self.assertAlmostEqual(doc.sections[0].page_width.mm, 210, delta=.1)
        self.assertAlmostEqual(doc.sections[0].left_margin.mm, 31.75, delta=.1)

    def test_template_preserves_reference_styles_and_theme(self):
        from sentinel_platform.core import project_root
        source = os.path.join(project_root(), "dicts", "report_templates", "ncc_event_source.docx")
        template = self.repo.collection("report_template").find_one({"_id": ncc.TEMPLATE_ID})
        with zipfile.ZipFile(source) as before, zipfile.ZipFile(os.path.join(self.root, template["template_path"])) as after:
            for part in before.namelist():
                if part.startswith(("word/styles", "word/theme/", "word/header", "word/footer")):
                    self.assertEqual(before.read(part), after.read(part), part)

    def test_tall_image_all_rows_fit_width_and_height(self):
        shot = rt.save_finding_shot("fixture", png((800, 3600)))
        self.assertTrue(shot.get("ok"))
        report = self.generate()
        doc = Document(os.path.join(self.root, report["docx_path"]))
        self.assertGreater(len(doc.inline_shapes), 1)
        section = doc.sections[0]
        for shape in doc.inline_shapes:
            self.assertLessEqual(shape.width, section.page_width - section.left_margin - section.right_margin)
            self.assertLessEqual(shape.height, section.page_height - section.top_margin - section.bottom_margin)
        with zipfile.ZipFile(os.path.join(self.root, report["docx_path"])) as z:
            heights = [Image.open(io.BytesIO(z.read(n))).height for n in z.namelist() if n.startswith("word/media/")]
            # Identical fragments can be de-duplicated in the ZIP; count their drawing references.
            rels = {r.rId: r.target_part.blob for r in doc.part.rels.values() if "image" in r.reltype}
            height = sum(Image.open(io.BytesIO(rels[shape._inline.graphic.graphicData.pic.blipFill.blip.embed])).height
                         for shape in doc.inline_shapes)
            self.assertEqual(height, 3600)

    def test_saved_edit_rebuilds_export_and_keeps_previous_file(self):
        shot = rt.save_finding_shot("fixture", png())
        report = copy.deepcopy(self.generate())
        old = os.path.join(self.root, report["docx_path"])
        data = copy.deepcopy(report["report_data"])
        data["report"]["title"] = "人工修订标题"
        finding = data["findings"][0]
        finding["root_cause"] = "人工确认的原因"
        finding.update(severity="low", severity_cn="低危")
        finding["reproduction_steps"] = [{"text": "先检查本地样例"}]
        finding["screenshots"][0].update(section="steps", step=1, caption="步骤一对应截图")
        result = asset_intel.update_report(report["_id"], title="人工修订标题", report_data=data)
        self.assertTrue(result.get("ok"), result)
        updated = self.repo.collection("pentest_report").find_one({"_id": report["_id"]})
        self.assertNotEqual(report["docx_path"], updated["docx_path"])
        self.assertEqual(updated["max_severity"], "low")
        self.assertEqual(self.repo.collection("intel_finding").find_one({"_id": "fixture"})["severity"], "high")
        self.assertTrue(os.path.isfile(old))
        with patch("sentinel_platform.core.template_dir", return_value=self.root):
            result = asset_intel.export_report_docx(report["_id"], collection="pentest_report")
        text = "\n".join(p.text for p in Document(io.BytesIO(result["data"])).paragraphs)
        self.assertIn("人工修订标题", text)
        self.assertIn("人工确认的原因", text)
        self.assertLess(text.index("步骤一对应截图"), text.index("4.证明材料"))

    def test_missing_image_fails_without_overwriting_saved_report(self):
        report = copy.deepcopy(self.generate())
        data = copy.deepcopy(report["report_data"])
        data["findings"][0]["screenshots"] = [{"name": "missing.png", "caption": "", "section": "evidence", "step": 1}]
        result = asset_intel.update_report(report["_id"], report_data=data)
        self.assertIn("error", result)
        saved = self.repo.collection("pentest_report").find_one({"_id": report["_id"]})
        self.assertEqual(saved["docx_path"], report["docx_path"])
        self.assertEqual(saved["report_data"], report["report_data"])

    def test_editor_cannot_replace_finding_or_load_arbitrary_file(self):
        report = self.generate()
        for mutation in ("identity", "path"):
            data = copy.deepcopy(report["report_data"])
            if mutation == "identity": data["findings"][0]["finding_id"] = "other"
            else: data["findings"][0]["screenshots"] = [{"name": "../other.png"}]
            self.assertIn("error", asset_intel.update_report(report["_id"], report_data=data))

    def test_template_markdown_edit_is_not_silently_ignored(self):
        report = self.generate()
        self.assertIn("error", asset_intel.update_report(report["_id"], content="invisible edit"))

    def test_regeneration_cannot_overwrite_manual_edits(self):
        report = self.generate()
        self.assertTrue(asset_intel.update_report(report["_id"], title="保留人工标题").get("ok"))
        result = rt.generate_from_template(ncc.TEMPLATE_ID, "finding", "fixture")
        self.assertIn("error", result)
        self.assertEqual(self.repo.collection("pentest_report").find_one({"_id": report["_id"]})["title"], "保留人工标题")

    def test_legacy_report_with_missing_source_cannot_be_emptied(self):
        self.repo.collection("pentest_report").insert_one({"_id": "legacy", "gen_mode": "template",
            "template_id": ncc.TEMPLATE_ID, "report_type": "finding", "vuln_index": ["missing"],
            "docx_path": "existing.docx", "title": "原报告"})
        self.assertIn("error", asset_intel.update_report("legacy", title="不应覆盖"))
        self.assertEqual(self.repo.collection("pentest_report").find_one({"_id": "legacy"})["title"], "原报告")

    def test_legacy_template_bound_images_are_included_once(self):
        legacy = rt.save_manual_shot(ncc.TEMPLATE_ID, "fixture", png(), "png")
        self.assertTrue(legacy.get("ok"), legacy)
        # The minimal fake repository does not implement Mongo's $push.
        self.repo.collection("report_template").find_one({"_id": ncc.TEMPLATE_ID})["manual_shots"] = {"fixture": [legacy["path"]]}
        report = self.generate()
        self.assertEqual(len(report["report_data"]["findings"][0]["screenshots"]), 1)
        self.generate()
        self.assertEqual(len(rt.list_finding_shots("fixture")), 1)


if __name__ == "__main__":
    unittest.main()
