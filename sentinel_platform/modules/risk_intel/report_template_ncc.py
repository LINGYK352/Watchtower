"""NCC event report: deterministic template built from the retained NCC reference."""
from __future__ import annotations

import copy
import os

TEMPLATE_ID = "ncc_event_finding_v1"
VERSION = 3


def build_template(source_path: str, output_path: str) -> None:
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    doc = Document(source_path)
    originals = list(doc.paragraphs)
    if len(originals) < 68 or "漏洞报告标题" not in originals[0].text:
        raise ValueError("NCC 原始模板结构不匹配")
    samples = {"title": copy.deepcopy(originals[0]._p),
               "heading": copy.deepcopy(originals[16]._p),
               "label": copy.deepcopy(originals[17]._p),
               "body": copy.deepcopy(originals[18]._p)}
    for child in list(doc._element.body):
        if not child.tag.endswith("}sectPr"):
            doc._element.body.remove(child)
    numbering = doc.part.numbering_part.element
    original_num = originals[17]._p.pPr.numPr.numId.val
    abstract = numbering.xpath('./w:num[@w:numId="{}"]/w:abstractNumId'.format(original_num))[0].get(qn("w:val"))
    def restart_numbering():
        num = numbering.add_num(int(abstract))
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        return num.numId
    number_id = restart_numbering()

    def paragraph(text, role="body", image=False, keep=False):
        nonlocal number_id
        if role == "heading":
            number_id = restart_numbering()
        element = copy.deepcopy(samples[role])
        for child in list(element):
            if not child.tag.endswith("}pPr"):
                element.remove(child)
        doc._element.body.insert(len(doc._element.body) - 1, element)
        p = Paragraph(element, doc._body)
        parts = text.split("{{", 1) if role == "title" else [text]
        run = p.add_run(parts[0])
        # Preserve the NCC font roles, including East Asian font selection.
        sample = originals[0 if role == "title" else 16 if role == "heading" else 17 if role == "label" else 18]
        if sample.runs and sample.runs[0]._r.rPr is not None:
            run._r.insert(0, copy.deepcopy(sample.runs[0]._r.rPr))
        if len(parts) == 2:
            value_run = p.add_run("{{" + parts[1])
            value_sample = originals[18].runs[0]
            if value_sample._r.rPr is not None:
                value_run._r.insert(0, copy.deepcopy(value_sample._r.rPr))
        if role == "label":
            p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = number_id
        p.paragraph_format.keep_with_next = keep or role in ("heading", "label")
        p.paragraph_format.keep_together = image
        p.paragraph_format.space_before = Pt(10 if role == "heading" else 0)
        p.paragraph_format.space_after = Pt(6)
        if image:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.left_indent = Pt(0)
            p.paragraph_format.right_indent = Pt(0)
            p.paragraph_format.first_line_indent = Pt(0)
            ind = p._p.pPr.find(qn("w:ind"))
            for key in ("firstLineChars", "leftChars", "rightChars", "hangingChars"):
                ind.attrib.pop(qn("w:" + key), None)
            p.paragraph_format.line_spacing = 1
        return p

    def field(label, variable, optional=False):
        if optional:
            paragraph("{%p if " + variable + " %}")
        paragraph(label, "label")
        value = paragraph("{{ " + variable + " | default('未提供', true) }}")
        if variable in ("item.poc", "item.evidence"):
            from docx.oxml import OxmlElement
            value.alignment = WD_ALIGN_PARAGRAPH.LEFT
            value.paragraph_format.first_line_indent = Pt(0)
            value.paragraph_format.line_spacing = 1.15
            value.paragraph_format.keep_together = True
            ind = value._p.pPr.find(qn("w:ind"))
            if ind is not None:
                ind.attrib.pop(qn("w:firstLineChars"), None)
            snap = OxmlElement("w:snapToGrid")
            snap.set(qn("w:val"), "0")
            value._p.pPr.append(snap)
        if optional:
            paragraph("{%p endif %}")

    def shots(variable):
        paragraph("{%p for shot in " + variable + " %}")
        paragraph("{%p if shot.caption %}")
        paragraph("{{ shot.caption }}", keep=True)
        paragraph("{%p endif %}")
        paragraph("{{ shot.image }}", image=True)
        paragraph("{%p endfor %}")

    paragraph("漏洞报告标题：{{ report_title }}", "title")
    paragraph("{%p for item in findings %}")
    paragraph("漏洞等级：{{ item.severity_cn }}", "title")
    paragraph("漏洞类型：{{ item.vuln_type }}", "title")
    paragraph("验证状态：{{ item.verification_status | default('待复核', true) }}", "title")
    paragraph("漏洞发现时间：{{ item.discovered_at | default('未提供', true) }}", "title")
    paragraph("受影响 IP：{{ item.affected_ip | default('未提供', true) }}", "title")
    paragraph("涉及单位：{{ report_unit | default('未提供', true) }}", "title")
    paragraph("单位全称：{{ report_unit | default('未提供', true) }}", "label")
    paragraph("受影响的平台/系统名称：{{ report_system | default('未提供', true) }}", "label")
    paragraph("访问地址：{{ item.target }}", "label")
    paragraph("版本信息：{{ item.version | default('未提供', true) }}", "label")
    paragraph("{%p if item.test_environment %}")
    paragraph("测试环境：{{ item.test_environment }}", "label")
    paragraph("{%p endif %}")
    paragraph("目标单位的域名 ICP 备案信息或归属证明：", "label")
    shots("item.icp_shots")
    paragraph("1.漏洞事件背景", "heading")
    for label, key in (("事件发现场景：", "discovery_context"), ("受影响用户类型：", "affected_users"),
                       ("受影响数据类型：", "affected_data"), ("事件量化信息：", "quantification")):
        field(label, "item." + key, optional=key == "quantification")
    paragraph("2.漏洞详细描述", "heading")
    for label, key in (("漏洞触发路径：", "trigger_path"), ("漏洞产生原因：", "root_cause"), ("事件影响范围：", "impact")):
        field(label, "item." + key)
    paragraph("3.复现步骤（可附截图/视频链接）", "heading")
    paragraph("{%p for step in item.reproduction_steps %}")
    paragraph("步骤 {{ loop.index }}：", "label")
    paragraph("{{ step.text | default('未提供', true) }}")
    shots("step.shots")
    paragraph("{%p endfor %}")
    field("复现条件与说明：", "item.verify_method")
    paragraph("4.证明材料", "heading")
    field("请求 / POC：", "item.poc")
    field("关键响应与结果：", "item.evidence")
    shots("item.evidence_shots")
    paragraph("5.修复建议", "heading")
    for label, key in (("紧急修复建议：", "remediation"), ("长期加固方案：", "hardening"), ("修复验证建议：", "fix_validation")):
        field(label, "item." + key, optional=key in ("hardening", "fix_validation"))
    paragraph("6.备注", "heading")
    paragraph("{{ item.notes | default('未提供', true) }}")
    paragraph("{%p endfor %}")
    doc.save(output_path)
    # python-docx reserializes even untouched XML. Restore opaque template parts
    # byte-for-byte; only body slots and numbering instances are authored here.
    import tempfile
    from zipfile import ZipFile
    fd, temporary = tempfile.mkstemp(suffix=".docx", dir=os.path.dirname(os.path.abspath(output_path)))
    os.close(fd)
    try:
        with ZipFile(source_path) as original, ZipFile(output_path) as modified, ZipFile(temporary, "w") as result:
            for part in original.infolist():
                package = modified if part.filename in ("word/document.xml", "word/numbering.xml") else original
                result.writestr(part, package.read(part.filename))
        os.replace(temporary, output_path)
    finally:
        if os.path.exists(temporary):
            os.remove(temporary)


def seed(force=False):
    from sentinel_platform.core import project_root
    from .report_template import repo_coll, _now, template_dir
    coll = repo_coll()
    existing = coll.find_one({"_id": TEMPLATE_ID}) or {}
    if not force and existing.get("ncc_version", 0) >= VERSION and os.path.isfile(
            os.path.join(template_dir(), existing.get("template_path", ""))):
        return 0
    root = os.path.join(template_dir(), TEMPLATE_ID)
    source = os.path.join(project_root(), "dicts", "report_templates", "ncc_event_source.docx")
    if not os.path.isfile(source):
        return 0
    os.makedirs(root, exist_ok=True)
    path = os.path.join(root, "ncc_event_v{}.docx".format(VERSION))
    # Keep previous files, including manually imported v1 templates, for reference.
    build_template(source, path)
    from .report_template_builtin import _builtin_schema, _sample_ctx
    from docxtpl import DocxTemplate
    origin = os.path.join(root, "ncc_event_v{}_origin.docx".format(VERSION))
    preview = DocxTemplate(path)
    preview.render(_sample_ctx("finding"), autoescape=True)
    preview.save(origin)
    meta = {"name": existing.get("name") or "NCC 事件型漏洞报告（可编辑）", "status": "ready",
            "scope": "finding", "builtin": True, "ncc_version": VERSION,
            "template_path": os.path.relpath(path, template_dir()).replace("\\", "/"),
            "origin_path": os.path.relpath(origin, template_dir()).replace("\\", "/"),
            "source_filename": "NCC事件型漏洞报告模版.docx", "schema": _builtin_schema(),
            "update_date": _now(), "need_review": False}
    meta["schema"]["images"] = []
    coll.update_one({"_id": TEMPLATE_ID}, {"$set": meta, "$setOnInsert": {"save_date": _now()}}, upsert=True)
    return 1
