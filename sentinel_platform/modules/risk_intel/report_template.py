"""risk_intel/report_template —— 报告模板学习 + 确定性生成（省 token 核心）。

设计（像"AI 拆解数学题→变成可计算机校验的结构"）：
  ① 学习（一次性 LLM）：用户上传报告模板 docx → python-docx 读出结构清单 → LLM 判定每个
     段落/表格是【固定文案 / 可变字段 / 循环表格 / 图表位】+ 映射语义字段 → 程序按判定在
     **用户原始 docx** 里注入 docxtpl 占位符（{{}} / {%tr%}），保留全部样式/排版/字体 XML。
  ② 生成（确定性，零/极少 LLM）：选模板 + 数据源(task/session) → 从 intel_finding 组装 context
     → matplotlib 画图表 PNG → docxtpl 只换文字/插图 → 输出 docx。同模板生成 N 篇增量 token=0。

防变形：以用户原模板为底，docxtpl 填充只替换占位符文本，排版零漂移。
关键坑（docxtpl 硬限制）：jinja2 标签必须落**单个 run**；Word 常把文字拆多 run → 注入时
  先合并 paragraph 的 run 再替换（见 _inject_scalar）。

依赖（惰性 import，缺失降级返错不崩）：python-docx（读+注入）、docxtpl（填充）、
  matplotlib（图表，Agg headless）。均需 rebuild 镜像（见 requirements.txt / Dockerfile）。
无 ROLE，经 registry 字符串键由 asset_intel 门面 IntelServiceImpl 转调。
"""
from __future__ import annotations

import os
import re
import uuid
import json
import shutil
from typing import Any, Dict, List, Optional

from sentinel_platform.core import get_repo, get_logger, template_dir, image_dir
from sentinel_platform.contracts import Collections

logger = get_logger()

# 严重度英文→中文（报告展示用；对齐 vuln_center 的 severity 归一值）
_SEV_CN = {"critical": "严重", "high": "高危", "medium": "中危", "low": "低危", "info": "信息"}
_SEV_ORDER = ["critical", "high", "medium", "low", "info"]


def _now() -> str:
    import time
    return time.strftime("%Y-%m-%d %H:%M:%S")


def _oid(v: Any):
    try:
        from bson import ObjectId
        return ObjectId(v)
    except Exception:
        return v


def _tpl_root(tid: str) -> str:
    """某模板的磁盘目录：<template_dir>/<tid>/。"""
    return os.path.join(template_dir(), str(tid))


# ========== ① 读 docx 结构（带稳定 index 锚点，供 AI 定位 + 程序回填）==========

def _read_docx_structure(path: str) -> List[Dict[str, Any]]:
    """读 docx → 结构清单：每段落/表格带稳定 index（p{n} / t{n}）+ 文本。
    段落 idx=p{顺序号}；表格 idx=t{顺序号}，附表头与前几行样例（供 AI 判断是否循环表）。
    只读文本，不改文件。python-docx 缺失则抛异常（上层捕获返错）。"""
    from docx import Document  # 惰性 import（需 rebuild 镜像装 python-docx）
    doc = Document(path)
    items: List[Dict[str, Any]] = []
    # document.paragraphs 与 tables 是文档级顺序独立的两个列表；这里各自编号，注入时按 kind+序号定位。
    for pi, p in enumerate(doc.paragraphs):
        txt = (p.text or "").strip()
        if txt:  # 空段落不喂 AI（省 token），注入时也不动
            items.append({"idx": "p{}".format(pi), "kind": "paragraph",
                          "style": (p.style.name if p.style else ""), "text": txt[:300]})
    for ti, tbl in enumerate(doc.tables):
        rows = tbl.rows
        header = [c.text.strip() for c in rows[0].cells] if rows else []
        sample = []
        for r in rows[1:3]:  # 取前两条数据行做样例
            sample.append([c.text.strip() for c in r.cells])
        items.append({"idx": "t{}".format(ti), "kind": "table",
                      "rows": len(rows), "header": header, "sample_rows": sample})
    return items


# ========== ② LLM 学习：结构清单 → schema（tool-use 强约束结构化输出）==========

# 语义字段字典（AI 把可变位映射到这些 key，生成时程序据此从 finding/stat/report 取值）
_SEMANTIC_FIELDS = {
    "report.title": "报告标题", "report.unit": "受测单位", "report.period": "测试时间",
    "report.system_name": "系统名称", "report.summary": "概述/执行摘要",
    "stat.total": "漏洞总数", "stat.critical": "严重数", "stat.high": "高危数",
    "stat.medium": "中危数", "stat.low": "低危数", "stat.info": "信息数",
    "finding.vuln_type": "漏洞名称", "finding.target": "目标/位置",
    "finding.severity_cn": "危害等级(中文)", "finding.impact": "危害描述",
    "finding.poc": "复现/PoC", "finding.verify_method": "验证方式",
    "finding.cvss_score": "CVSS 评分", "finding.evidence": "证据",
    "chart.severity_dist": "严重度分布图", "chart.attack_chain": "攻击链图",
}

_LEARN_TOOL = {
    "name": "emit_template_schema",
    "description": "输出报告模板的结构判定：每个元素是固定文案/可变字段/循环表格/图表位，并映射语义字段。",
    "parameters": {
        "type": "object",
        "properties": {
            "scalars": {"type": "array", "description": "可变标量字段（填一个值的位置）",
                        "items": {"type": "object", "properties": {
                            "idx": {"type": "string", "description": "元素锚点，如 p5"},
                            "placeholder": {"type": "string", "description": "jinja2 占位符名(英文/下划线)"},
                            "semantic": {"type": "string", "description": "语义字段(见字段字典)或 custom"},
                            "sample_text": {"type": "string", "description": "该段落里应被替换的原示例值(精确子串)；给不到留空"}}}},
            "loops": {"type": "array", "description": "循环表格（每行一条漏洞）",
                      "items": {"type": "object", "properties": {
                          "idx": {"type": "string", "description": "表格锚点，如 t1"},
                          "data_row": {"type": "integer", "description": "数据样例行的行号(0基,通常1,表头后第一行)"},
                          "columns": {"type": "array", "items": {"type": "object", "properties": {
                              "col": {"type": "integer"}, "placeholder": {"type": "string"},
                              "semantic": {"type": "string"}}}}}}},
            "images": {"type": "array", "description": "图表位（插图）",
                       "items": {"type": "object", "properties": {
                           "idx": {"type": "string"}, "placeholder": {"type": "string"},
                           "semantic": {"type": "string", "description": "chart.severity_dist 或 chart.attack_chain"},
                           "sample_text": {"type": "string"}}}},
            "fixed_kept": {"type": "array", "items": {"type": "string"}, "description": "固定文案元素 idx(原样保留)"},
        },
        "required": ["scalars", "loops"],
    },
}


def _resolve_learn_provider(provider_id: str = "") -> Optional[Dict[str, Any]]:
    """解析模板学习用的 provider。
    provider_id 非空 → 按 id 取（可用才用，不可用降级）；空 → template_learn 场景 → 全局默认。
    返回可用 provider dict 或 None。"""
    from sentinel_platform.modules.ai_pentest import _llm
    pid = (provider_id or "").strip()
    if pid:
        try:
            from sentinel_platform.modules.ai_pentest import ai_config
            prov = ai_config.get_provider(pid, require_enabled=True)
            if prov and ai_config._provider_usable(prov):
                return prov
            logger.warning("report_template: 指定学习模型 %s 不可用，降级 template_learn 场景/全局默认", pid)
        except Exception as exc:
            logger.debug("report_template: 取指定学习模型 %s 失败：%s", pid, exc)
    prov = _llm.resolve_provider("template_learn")
    if not prov:
        # 场景没配就退全局默认（同 _report_llm_provider 语义）
        try:
            from sentinel_platform.contracts import get_registry
            svc = get_registry().get("ai_config_service")
            prov = svc.get_active_provider() if svc and hasattr(svc, "get_active_provider") else None
        except Exception:
            prov = None
    return prov if (prov and prov.get("api_key")) else None


def _learn_schema_via_llm(structure: List[Dict[str, Any]], provider_id: str = "",
                          feedback: str = "") -> Dict[str, Any]:
    """喂结构清单给 LLM，用 tool-use 强约束输出 schema。返回 {ok, schema, tokens, provider_id, provider_name, error?}。
    provider_id 非空按 id 取（不可用降级）；空走 template_learn 场景→全局默认。
    feedback 非空时作为额外约束并入（人在回路重学：用户指出上次学错哪、这次怎么改）。"""
    from sentinel_platform.modules.ai_pentest import _llm
    prov = _resolve_learn_provider(provider_id)
    if not prov or not prov.get("api_key"):
        return {"ok": False, "error": "无可用 AI 模型（请在 AI 配置启用 provider 或绑定 template_learn 场景）"}

    fields_doc = "\n".join("  {} = {}".format(k, v) for k, v in _SEMANTIC_FIELDS.items())
    sys_prompt = (
        "你是报告模板结构分析器。给你一份渗透报告模板的元素清单（每个元素带锚点 idx 和文本）。\n"
        "判断每个元素属于哪类，**只输出结构判定，绝不重写正文**：\n"
        "① 固定文案(fixed_kept)：章节标题、说明性文字、法律声明等每份报告都一样的→放 fixed_kept。\n"
        "② 可变标量(scalars)：单位名/日期/系统名/统计数字等每份报告不同的单值→给 placeholder(英文下划线)+semantic。\n"
        "   若该段里只有部分文字可变(如'测试时间：2024年3月'的日期部分)，把可变的原文精确填进 sample_text。\n"
        "③ 循环表格(loops)：漏洞清单这类每行一条记录、行数随报告变化的表格→标注数据样例行 data_row 和每列 semantic。\n"
        "④ 图表位(images)：需要插入饼图/攻击链图的位置→semantic 用 chart.severity_dist 或 chart.attack_chain。\n"
        "可用语义字段字典(semantic 尽量取这些，取不到用 custom)：\n" + fields_doc + "\n"
        "调用 emit_template_schema 返回结构。"
    )
    user = "模板元素清单（JSON）：\n" + json.dumps(structure, ensure_ascii=False)
    if (feedback or "").strip():
        # 人在回路重学：把用户对上一版的纠正意见作为额外约束（放 user 段末，权重高）
        user += ("\n\n【用户对上一版模板学习结果的纠正意见（请据此修正判定）】\n"
                 + feedback.strip())
    msgs = [{"role": "system", "content": sys_prompt}, {"role": "user", "content": user}]
    r = _llm.chat(prov, msgs, timeout=180.0, tools=[_LEARN_TOOL], scene="template_learn")
    if not r.get("ok"):
        return {"ok": False, "error": "AI 学习失败：{}".format(r.get("error") or "空返回")}
    # 取 tool_call 的结构化参数（_llm 已把 arguments 解析成 dict）
    schema = None
    for tc in (r.get("tool_calls") or []):
        if tc.get("name") == "emit_template_schema":
            schema = tc.get("arguments") or {}
            break
    if schema is None:
        # 模型没走 tool、只吐了文本 → 尝试从 content 里抠 JSON（宽容兜底）
        m = re.search(r"\{[\s\S]*\}", r.get("content") or "")
        if m:
            try:
                schema = json.loads(m.group(0))
            except Exception:
                schema = None
    if not isinstance(schema, dict):
        return {"ok": False, "error": "AI 未产出可用的模板结构（未调用 emit_template_schema）"}
    schema.setdefault("version", 1)
    schema.setdefault("scalars", [])
    schema.setdefault("loops", [])
    schema.setdefault("images", [])
    schema.setdefault("fixed_kept", [])
    return {"ok": True, "schema": schema, "tokens": int(r.get("tokens", 0) or 0),
            "provider_id": str(prov.get("_id", "")), "provider_name": prov.get("name", "") or prov.get("model", "")}


# ========== ③ 占位符注入（解决 docxtpl 标签跨 run 限制）==========

def _merge_paragraph_runs(paragraph) -> None:
    """把 paragraph 的所有 run 文本归并进第一个 run，清空其余 run 文本（保留 runs[0] 的 rPr 样式）。
    docxtpl/jinja2 标签必须落单个 run；Word 常把文字拆多 run（拼写检查/格式痕迹），
    合并后标签才不会被拆碎。空段落/无 run 直接跳过。"""
    runs = paragraph.runs
    if len(runs) <= 1:
        return
    full = "".join(r.text or "" for r in runs)
    runs[0].text = full
    for r in runs[1:]:
        r.text = ""


def _inject_scalar(paragraph, placeholder: str, sample: str) -> None:
    """在（已合并 run 的）段落里注入 {{ placeholder }}。
    有 sample→只替换该子串（保留段落其余固定文字，如"测试时间："前缀）；无 sample→整段文本换成占位符。"""
    _merge_paragraph_runs(paragraph)
    runs = paragraph.runs
    if not runs:
        # 无 run（罕见）→ 直接加一个
        paragraph.add_run("{{ %s }}" % placeholder)
        return
    tag = "{{ %s }}" % placeholder
    if sample and sample in runs[0].text:
        runs[0].text = runs[0].text.replace(sample, tag)
    else:
        runs[0].text = tag


# 循环列语义 → finding dict 字段名（对齐 _collect_findings / _mock_finding 产出的键）。
# 取不到映射时按列序兜底 col{n}，保证注入侧(item.X)与预览/生成侧提供的键一致，永不出现 undefined 变量。
_LOOP_FIELD_ALIAS = {
    "seq": "seq", "finding.seq": "seq",
    "finding.vuln_type": "vuln_type", "vuln_type": "vuln_type", "漏洞名称": "vuln_type",
    "finding.target": "target", "target": "target", "目标": "target", "位置": "target",
    "finding.severity_cn": "severity_cn", "severity_cn": "severity_cn", "危害等级": "severity_cn", "等级": "severity_cn",
    "finding.impact": "impact", "impact": "impact", "危害": "impact",
    "finding.poc": "poc", "poc": "poc", "复现": "poc",
    "finding.verify_method": "verify_method", "verify_method": "verify_method", "验证方式": "verify_method",
    "finding.cvss_score": "cvss_score", "cvss_score": "cvss_score",
    "finding.evidence": "evidence", "evidence": "evidence", "证据": "evidence",
}


def _loop_field_name(col: Dict[str, Any]) -> str:
    """从循环列的 semantic 派生 finding 字段名（不信 AI 自由填的 placeholder 变量名）。
    优先 semantic 映射；取不到用 placeholder 末段；再退 col{序号}。返回 jinja2 安全的字段名。"""
    sem = (col.get("semantic") or "").strip()
    if sem in _LOOP_FIELD_ALIAS:
        return _LOOP_FIELD_ALIAS[sem]
    # semantic 形如 finding.xxx 未命中别名 → 取末段
    if "." in sem:
        tail = sem.split(".")[-1].strip()
        if tail:
            return re.sub(r"[^0-9a-zA-Z_]", "_", tail)
    # placeholder 末段兜底（如 item.vuln_type / row.x → vuln_type / x）
    ph = (col.get("placeholder") or "").strip()
    if ph:
        tail = ph.split(".")[-1].strip()
        if tail:
            return re.sub(r"[^0-9a-zA-Z_]", "_", tail)
    return "col{}".format(col.get("col", 0))


def _inject_placeholders(doc, schema: Dict[str, Any]) -> None:
    """按 schema 在 python-docx Document 上注入占位符。就地改，调用方负责 save。"""
    paras = doc.paragraphs
    tables = doc.tables

    def _para_by_idx(idx: str):
        # idx 形如 p12
        try:
            return paras[int(idx[1:])]
        except (ValueError, IndexError):
            return None

    # 标量
    for s in schema.get("scalars", []):
        idx = str(s.get("idx", ""))
        ph = (s.get("placeholder") or "").strip()
        if not idx.startswith("p") or not ph:
            continue
        p = _para_by_idx(idx)
        if p is not None:
            _inject_scalar(p, ph, (s.get("sample_text") or "").strip())

    # 图表位（当占位符处理，值在生成时以 InlineImage 传入）
    for im in schema.get("images", []):
        idx = str(im.get("idx", ""))
        ph = (im.get("placeholder") or "").strip()
        if not idx.startswith("p") or not ph:
            continue
        p = _para_by_idx(idx)
        if p is not None:
            _inject_scalar(p, ph, (im.get("sample_text") or "").strip())

    # 循环表格：数据行注入 {%tr for %}/{%tr endfor %}（行循环，标签独占 run 不与数据混排）
    for lp in schema.get("loops", []):
        idx = str(lp.get("idx", ""))
        if not idx.startswith("t"):
            continue
        try:
            tbl = tables[int(idx[1:])]
        except (ValueError, IndexError):
            continue
        data_row = int(lp.get("data_row", 1) or 1)
        if data_row >= len(tbl.rows):
            continue
        row = tbl.rows[data_row]
        cells = row.cells
        if not cells:
            continue
        # 各列注入占位符：**强制循环变量 item + 字段名由 semantic 派生**（不信 AI 自由填的 placeholder）。
        # 治弱模型把列占位符填成 row.x / 其它变量名 → 与循环变量 item 对不上 → 渲染报 'row' is undefined 失败
        # （E2E 实测 DeepSeek 重学踩到）。字段名对齐 _collect_findings / _mock_finding 的 finding dict 键。
        for col in lp.get("columns", []):
            ci = int(col.get("col", -1))
            if ci < 0 or ci >= len(cells):
                continue
            _inject_scalar(cells[ci].paragraphs[0], "item.{}".format(_loop_field_name(col)), "")
        # docxtpl 的 {%tr%} 是**整行**操作：for 与 endfor 必须各占**独立一行**（同一行首尾放会
        # 报 "unknown tag endfor"，实测踩坑）。故在数据行前插 for 标记行、后插 endfor 标记行，
        # docxtpl 渲染时删掉这两行，仅循环中间数据行。
        _insert_tag_row(tbl, data_row, "{%tr for item in findings %}", before=True)
        # 前面插了一行 → 数据行下移一位，endfor 插到数据行(now data_row+1)之后
        _insert_tag_row(tbl, data_row + 1, "{%tr endfor %}", before=False)


def _insert_tag_row(tbl, ref_row_idx: int, tag: str, before: bool) -> None:
    """在表格 ref_row_idx 行的前/后插入一整行，首单元格放 docxtpl 行循环标签({%tr%})。
    用 add_row 造出列数一致的空行，再用 lxml 移到目标位置。"""
    new_row = tbl.add_row()   # 追加到表尾（列数=表列数）
    new_row.cells[0].text = tag
    tr = new_row._tr
    tr.getparent().remove(tr)   # 从表尾摘出
    ref_tr = tbl.rows[ref_row_idx]._tr
    if before:
        ref_tr.addprevious(tr)
    else:
        ref_tr.addnext(tr)


# ========== 学习入口（上传后调用）==========

def _run_learn_cycle(tid: str, origin_dst: str, root: str, provider_id: str = "",
                     feedback: str = "", need_review: bool = False,
                     incr_refine: bool = False) -> Dict[str, Any]:
    """学习循环核心（learn 首学 + refine 重学共用）：读结构→LLM 学→注入→自检→落库。
    origin_dst 必须已存在（learn 归档后 / refine 复用历史）。失败落 status=failed + learn_error。
    need_review=True → 成功落 status=review（待人工确认）；否则 ready（可直接生成）。"""
    repo = repo_coll()

    def _prog(phase: str, pct: int) -> None:
        """更新学习进度（异步学习期间前端轮询列表读 learn_phase/learn_progress 显进度条）。"""
        try:
            repo.update_one({"_id": _oid(tid)},
                            {"$set": {"learn_phase": phase, "learn_progress": int(pct),
                                      "update_date": _now()}})
        except Exception:
            pass

    def _fail(msg: str) -> Dict[str, Any]:
        repo.update_one({"_id": _oid(tid)},
                        {"$set": {"status": "failed", "learn_error": msg[:500],
                                  "learn_phase": "失败", "learn_progress": 100, "update_date": _now()}})
        return {"ok": False, "template_id": tid, "status": "failed", "error": msg}

    # ① 读结构
    _prog("读取文档结构", 10)
    try:
        structure = _read_docx_structure(origin_dst)
    except Exception as exc:
        return _fail("读取 docx 失败（可能非 .docx 或组件未装）：{}".format(exc))
    if not structure:
        return _fail("模板内容为空，无可学习结构")

    # ② LLM 学 schema（provider_id 可选，feedback 用于人在回路重学）—— 最慢一步
    _prog("AI 学习结构中", 30)
    learned = _learn_schema_via_llm(structure, provider_id=provider_id, feedback=feedback)
    if not learned.get("ok"):
        return _fail(learned.get("error", "学习失败"))
    schema = learned["schema"]
    tokens = learned.get("tokens", 0)

    # ②.5 空 schema 兜底（治"效果不理想"：弱模型对复杂真实报告常产出全空 schema）：
    # scalars+loops+images 全空 = 没学到任何可变结构/循环表 → 落 failed 而非 review，避免空模板伪装待确认。
    # fixed_kept 不算有效结构（全固定=没识别出可填充位，等于没学会）。
    if not (schema.get("scalars") or schema.get("loops") or schema.get("images")):
        return _fail("未学到任何可变字段/循环表/图表位（schema 为空）——可能模型能力不足或文档结构复杂，"
                     "请换更强模型（如 Claude 高配）重学。")

    # ③ 注入占位符 → template.docx
    _prog("注入占位符", 75)
    template_dst = os.path.join(root, "template.docx")
    try:
        from docx import Document
        d = Document(origin_dst)
        _inject_placeholders(d, schema)
        d.save(template_dst)
    except Exception as exc:
        return _fail("占位符注入失败：{}".format(exc))

    # ④ 自检：docxtpl + mock context 试渲染一次，成功才落 review/ready（带病模板挡在学习阶段）
    _prog("自检渲染", 90)
    ok, err = _self_check_render(template_dst, schema)
    if not ok:
        return _fail("模板自检渲染失败（可能标签跨 run/结构学错）：{}".format(err))

    final_status = "review" if need_review else "ready"
    upd = {"status": final_status, "schema": schema, "learn_tokens": tokens, "learn_error": "",
           "learn_phase": "完成", "learn_progress": 100,
           "learn_provider_id": learned.get("provider_id", ""),
           "learn_provider_name": learned.get("provider_name", ""),
           "template_path": os.path.relpath(template_dst, template_dir()).replace("\\", "/"),
           "origin_path": os.path.relpath(origin_dst, template_dir()).replace("\\", "/"),
           "update_date": _now()}
    op: Dict[str, Any] = {"$set": upd}
    if incr_refine:
        op["$inc"] = {"refine_count": 1}
    repo.update_one({"_id": _oid(tid)}, op)
    return {"ok": True, "template_id": tid, "status": final_status, "tokens": tokens,
            "provider_name": learned.get("provider_name", "")}


def learn_template(name: str, origin_path: str, source_filename: str = "",
                   created_by: str = "", provider_id: str = "",
                   need_review: bool = True, sync: bool = False) -> Dict[str, Any]:
    """学习一份上传的模板 docx。**默认异步**（sync=False）：同步建 learning 文档 + 归档原件后立即返回，
    LLM 学习循环在后台 daemon 线程跑（上传端点秒回，前端列表显示"学习中+进度"，学习期间可用其他功能）。
    进度经 mongo learn_phase/learn_progress 字段（前端轮询列表）。
    provider_id 空=走 template_learn 场景/全局默认；need_review 默认 True（学成落 review 待人工确认）。
    sync=True 用于单测/scheduler 兜底（同步跑完返回终态）。
    **归档必须同步**（端点 finally 删临时件），学习循环才可异步。"""
    repo = repo_coll()
    now = _now()
    # 先建 learning 文档拿到 id（后续文件按 id 归档）
    doc = {"name": (name or "未命名模板").strip(), "status": "learning",
           "source_filename": source_filename or os.path.basename(origin_path),
           "schema": {}, "learn_tokens": 0, "learn_error": "", "refine_count": 0,
           "learn_phase": "排队中", "learn_progress": 0,
           "learn_provider_id": provider_id or "", "learn_provider_name": "",
           "need_review": bool(need_review),
           "created_by": created_by or "", "save_date": now, "update_date": now}
    tid = str(repo.insert_one(doc).inserted_id)
    root = _tpl_root(tid)
    os.makedirs(root, exist_ok=True)

    try:
        # 原件归档到模板目录（端点存的临时件 → origin.docx）——**同步做**，因端点 finally 会删临时件
        origin_dst = os.path.join(root, "origin.docx")
        if os.path.abspath(origin_path) != os.path.abspath(origin_dst):
            shutil.copyfile(origin_path, origin_dst)
    except Exception as exc:
        repo.update_one({"_id": _oid(tid)},
                        {"$set": {"status": "failed", "learn_error": "原件归档失败：{}".format(exc)[:500],
                                  "learn_phase": "失败", "learn_progress": 100, "update_date": _now()}})
        return {"ok": False, "template_id": tid, "status": "failed", "error": "原件归档失败：{}".format(exc)}

    if sync:
        return _run_learn_cycle(tid, origin_dst, root, provider_id=provider_id, need_review=need_review)

    # 异步：后台 daemon 线程跑学习循环。学习只写 template_dir/<tid>/ + mongo doc，不改任何 .py，
    # 不触发 gunicorn --reload（故 worker 内线程安全，见 plan 探索结论）；worker 崩则 scheduler tick 兜底重跑。
    import threading

    def _bg():
        try:
            _run_learn_cycle(tid, origin_dst, root, provider_id=provider_id, need_review=need_review)
        except Exception as exc:
            logger.warning("report_template async learn %s crashed: %s", tid, exc)
            try:
                repo.update_one({"_id": _oid(tid)},
                                {"$set": {"status": "failed", "learn_error": "学习异常：{}".format(exc)[:500],
                                          "learn_phase": "失败", "learn_progress": 100, "update_date": _now()}})
            except Exception:
                pass

    threading.Thread(target=_bg, name="tpl-learn-{}".format(tid[:8]), daemon=True).start()
    return {"ok": True, "template_id": tid, "status": "learning", "async": True}


def refine_template(template_id: str, provider_id: str = "", feedback: str = "") -> Dict[str, Any]:
    """人在回路重学（re-learn）：复用已归档的 origin.docx 重新学 schema，覆盖 template.docx，仍落 review 态。
    可换更强模型（provider_id）+ 带用户纠正意见（feedback）。反复直到用户满意。返回 {ok, status, ...}。"""
    tid = (template_id or "").strip()
    if not tid:
        return {"error": "template_id 必填"}
    d = repo_coll().find_one({"_id": _oid(tid)})
    if not d:
        return {"error": "模板不存在"}
    root = _tpl_root(tid)
    origin_dst = os.path.join(root, "origin.docx")
    if not os.path.isfile(origin_dst):
        return {"error": "原样本文件缺失，无法重学（请重新上传）"}
    # 重学固定进 review 态（既然主动 refine，必然要再复核确认）
    r = _run_learn_cycle(tid, origin_dst, root, provider_id=provider_id,
                         feedback=feedback, need_review=True, incr_refine=True)
    if not r.get("ok"):
        return {"error": r.get("error", "重学失败"), "status": r.get("status", "failed")}
    return r


def confirm_template(template_id: str) -> Dict[str, Any]:
    """确认定稿：review → ready（之后才可用于生成）。非 review 态拒绝（防误操作）。"""
    tid = (template_id or "").strip()
    if not tid:
        return {"error": "template_id 必填"}
    d = repo_coll().find_one({"_id": _oid(tid)})
    if not d:
        return {"error": "模板不存在"}
    st = d.get("status")
    if st == "ready":
        return {"ok": True, "template_id": tid, "status": "ready", "note": "已是定稿态"}
    if st != "review":
        return {"error": "仅待确认(review)态可定稿，当前 status={}".format(st)}
    repo_coll().update_one({"_id": _oid(tid)},
                           {"$set": {"status": "ready", "update_date": _now()}})
    return {"ok": True, "template_id": tid, "status": "ready"}


def template_diff(template_id: str) -> Dict[str, Any]:
    """在线并排对比数据（不渲染 docx）：origin 每段原文 vs schema 对该 idx 的判定。
    返回 {ok, name, status, rows:[{idx, kind, text, verdict, semantic}], summary}。供前端复核抽屉左右并排。"""
    tid = (template_id or "").strip()
    if not tid:
        return {"error": "template_id 必填"}
    d = repo_coll().find_one({"_id": _oid(tid)})
    if not d:
        return {"error": "模板不存在"}
    origin_dst = os.path.join(_tpl_root(tid), "origin.docx")
    if not os.path.isfile(origin_dst):
        return {"error": "原样本文件缺失"}
    try:
        structure = _read_docx_structure(origin_dst)
    except Exception as exc:
        return {"error": "读取原样本失败：{}".format(exc)}
    schema = d.get("schema") or {}
    # idx → 判定映射（scalar/loop/image/fixed）
    verdict_map: Dict[str, Dict[str, str]] = {}
    for s in schema.get("scalars", []):
        i = s.get("idx") or s.get("anchor") or ""
        if i:
            verdict_map[i] = {"verdict": "可变字段", "semantic": s.get("semantic") or s.get("placeholder") or ""}
    for lp in schema.get("loops", []):
        i = lp.get("idx") or lp.get("anchor") or ""
        if i:
            cols = lp.get("columns") or lp.get("cols") or []
            verdict_map[i] = {"verdict": "循环表格", "semantic": ",".join(
                [str(c.get("semantic") or c) if isinstance(c, dict) else str(c) for c in cols][:8])}
    for im in schema.get("images", []):
        i = im.get("idx") or im.get("anchor") or ""
        if i:
            verdict_map[i] = {"verdict": "图表位", "semantic": im.get("semantic") or ""}
    fixed_idx = set(schema.get("fixed_kept", []) or [])
    rows = []
    for el in structure:
        idx = el.get("idx", "")
        if idx in verdict_map:
            v = verdict_map[idx]
        elif idx in fixed_idx:
            v = {"verdict": "固定文案", "semantic": ""}
        else:
            v = {"verdict": "未识别", "semantic": ""}   # 既没进可变/循环/图，也没进 fixed → 提醒用户
        rows.append({"idx": idx, "kind": el.get("kind", ""),
                     "text": (el.get("text") or "") if el.get("kind") == "paragraph"
                             else "表格：{} 行，表头 {}".format(el.get("rows", 0), el.get("header", [])),
                     "verdict": v["verdict"], "semantic": v["semantic"]})
    summary = {"total": len(rows),
               "scalar": len(schema.get("scalars", [])), "loop": len(schema.get("loops", [])),
               "image": len(schema.get("images", [])), "fixed": len(fixed_idx),
               "unrecognized": sum(1 for r in rows if r["verdict"] == "未识别")}
    return {"ok": True, "name": d.get("name", ""), "status": d.get("status", ""),
            "rows": rows, "summary": summary}


def get_template_docx(template_id: str, which: str = "origin") -> Dict[str, Any]:
    """取模板的原始 docx 绝对路径供在线对比渲染：which=origin（原报告）/ template（打标记的模板）。
    前端用 mammoth.js 把这两份 docx 转 HTML 左右并排（左=真实报告排版，右=可变位显 {{占位符}}/循环标记）。
    返回 {ok, path(绝对), filename} 或 {error}。端点侧做 realpath 防遍历后下发字节。"""
    tid = (template_id or "").strip()
    if not tid:
        return {"error": "template_id 必填"}
    d = repo_coll().find_one({"_id": _oid(tid)})
    if not d:
        return {"error": "模板不存在"}
    root = _tpl_root(tid)
    fname = "template.docx" if which == "template" else "origin.docx"
    path = os.path.join(root, fname)
    if not os.path.isfile(path):
        return {"error": "{} 文件缺失（学习可能未完成或失败）".format(fname)}
    return {"ok": True, "path": path, "filename": fname}


def run_pending_learn(stale_seconds: int = 600) -> Dict[str, Any]:
    """scheduler 兜底：捡 status=learning 且心跳(update_date)超时的模板（worker 崩溃线程死），同步重跑学习。
    正常异步学习几分钟内完成；超 stale_seconds 仍 learning = 线程已死 → 重跑。返回 {reclaimed, ids}。"""
    import time as _t
    reclaimed, ids = 0, []
    cutoff = _t.strftime("%Y-%m-%d %H:%M:%S", _t.localtime(_t.time() - max(60, stale_seconds)))
    stale = list(repo_coll().find({"status": "learning", "update_date": {"$lt": cutoff}}))
    for d in stale:
        tid = str(d.get("_id"))
        root = _tpl_root(tid)
        origin_dst = os.path.join(root, "origin.docx")
        if not os.path.isfile(origin_dst):
            # 原件都没了，无法重跑 → 落 failed 清僵尸
            repo_coll().update_one({"_id": _oid(tid)},
                                   {"$set": {"status": "failed", "learn_error": "学习中断且原样本缺失，请重新上传",
                                             "learn_phase": "失败", "learn_progress": 100, "update_date": _now()}})
            continue
        _run_learn_cycle(tid, origin_dst, root, provider_id=d.get("learn_provider_id", "") or "",
                         need_review=bool(d.get("need_review", True)))
        reclaimed += 1
        ids.append(tid)
    return {"reclaimed": reclaimed, "ids": ids}


def _self_check_render(template_path: str, schema: Dict[str, Any]) -> tuple:
    """用 mock context 试渲染模板，成功返 (True,'')，失败返 (False, err)。"""
    try:
        from docxtpl import DocxTemplate
        tpl = DocxTemplate(template_path)
        ctx: Dict[str, Any] = {"findings": [_mock_finding()]}
        for s in schema.get("scalars", []):
            ph = (s.get("placeholder") or "").strip()
            if ph:
                ctx.setdefault(ph, "样例")
        for im in schema.get("images", []):
            ph = (im.get("placeholder") or "").strip()
            if ph:
                ctx.setdefault(ph, "")  # 自检不插真图，空串即可验证标签合法
        tpl.render(ctx, autoescape=True)   # 与真实生成一致（见 generate_from_template 注释）
        return True, ""
    except Exception as exc:
        return False, str(exc)


def _mock_finding() -> Dict[str, Any]:
    return {"seq": 1, "vuln_type": "样例漏洞", "target": "https://example.com",
            "severity_cn": "高危", "impact": "样例危害", "poc": "样例PoC",
            "verify_method": "样例验证", "cvss_score": "7.5", "evidence": "样例证据"}


def repo_coll():
    return get_repo().collection(Collections.REPORT_TEMPLATE)


# ========== ④ 生成链路（确定性填充，零/极少 LLM）==========

def _collect_findings(source: str, source_id: str) -> List[Dict[str, Any]]:
    """按数据源取 finding 列表：session→按 session_id；task→按该任务下会话报告的 vuln_index。
    映射成模板可用的扁平 dict（语义字段对齐 _SEMANTIC_FIELDS 的 finding.*）。"""
    repo = get_repo()
    fcoll = repo.collection(Collections.INTEL_FINDING)
    raws: List[Dict[str, Any]] = []
    fid_system: Dict[str, str] = {}   # finding_id → 所属系统名（任务级分组用；session 源不填）
    if source == "finding":   # 单漏洞级：source_id = intel_finding 的 _id，取该单条
        try:
            one = fcoll.find_one({"_id": _oid(source_id)})
        except Exception:
            one = None
        raws = [one] if one else []
    elif source == "session":
        raws = list(fcoll.find({"session_id": source_id, "duplicate_of": None}))
    else:  # task：复刻 _task_findings_summary 的 vuln_index 归集
        reports = list(repo.collection(Collections.INTEL_REPORT).find(
            {"source_task_id": str(source_id), "report_type": {"$ne": "task"}}))
        fids = set()
        for r in reports:
            sysname = (r.get("system_name") or "").strip()
            for fid in (r.get("vuln_index") or []):
                sfid = str(fid)
                fids.add(sfid)
                # 一条 finding 可能被多份会话报告引用；首个非空 system_name 胜出（稳定归属）
                if sysname and not fid_system.get(sfid):
                    fid_system[sfid] = sysname
        for fid in fids:
            try:
                f = fcoll.find_one({"_id": _oid(fid)})
            except Exception:
                f = None
            if f:
                raws.append(f)
    out = []
    seq = 0
    seen = set()
    for f in raws:
        alias_seen = set()
        while f.get("duplicate_of") and str(f.get("_id")) not in alias_seen:
            alias_seen.add(str(f.get("_id")))
            replacement = fcoll.find_one({"_id": _oid(f["duplicate_of"])})
            if not replacement:
                break
            if fid_system.get(str(f.get("_id"))):
                fid_system.setdefault(str(replacement["_id"]), fid_system[str(f["_id"])])
            f = replacement
        if f.get("duplicate_of") or str(f.get("_id")) in seen:
            continue
        seen.add(str(f.get("_id")))
        # 排除人工标记的误报（v1.21.157-48 item6）：误报漏洞不进报告。降级已自动生效（读 severity=降后值）。
        if (f.get("handle_status") or "") == "false_positive":
            continue
        seq += 1
        i = seq
        sv = (f.get("severity") or "").lower()
        fid = str(f.get("_id", ""))
        out.append({
            "seq": i, "finding_id": fid,
            "vuln_type": f.get("vuln_type") or f.get("name") or "未命名",
            "target": f.get("target") or f.get("asset_key") or "",
            "severity": sv, "severity_cn": _SEV_CN.get(sv, sv or "-"),
            "impact": f.get("impact") or "", "poc": f.get("poc") or "",
            "verify_method": ("待验证；" if not f.get("verified") else "") + (f.get("verify_method") or ""),
            "verification_status": "已验证" if f.get("verified") else "待验证线索",
            "poc_notes": f.get("poc_notes") or "",
            "cvss_score": f.get("cvss_score") or "",
            "evidence": _evidence_text(f.get("evidence")),
            "session_id": f.get("session_id") or "",
            "asset_key": f.get("asset_key") or "",
            # 任务级按系统分组用：从产出该 finding 的会话报告继承 system_name（session 源留空，由 meta 补）
            "system_name": fid_system.get(fid, "") if source == "task" else "",
        })
        from .report_document import FINDING_FIELDS
        for key in FINDING_FIELDS:
            out[-1].setdefault(key, str(f.get(key) or ""))
        out[-1]["discovered_at"] = str(f.get("discovered_at") or f.get("save_date") or "")[:10]
        out[-1]["reproduction_steps"] = f.get("reproduction_steps") or []
        out[-1]["evidence_records"] = [e for e in (f.get("evidence") or []) if isinstance(e, dict)] if isinstance(f.get("evidence"), list) else []
    return out


def _group_findings_by_system(source: str, findings: List[Dict[str, Any]],
                              meta_system: str = "") -> List[Dict[str, Any]]:
    """把扁平 findings 按「所属系统」分组，供任务级模板多一级「分系统」目录渲染。
    - session 源：单系统 → 单组（组名取会话 system_name，缺省"目标系统"）。
    - task 源：按每条 finding 的 system_name 分桶（缺省"未分组系统"），组内漏洞序号从 1 重排。
    返回 [{name, findings(已重排 seq), stat}]，stat=该组严重度统计（供模板 {{ sys.stat.high }} 等）。"""
    from collections import OrderedDict

    def _reseq(fs: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        for i, f in enumerate(fs, 1):
            f = dict(f)
            f["seq"] = i
            fs[i - 1] = f
        return fs

    if source != "task":
        fs = _reseq([dict(f) for f in findings])
        return [{"name": meta_system or "目标系统", "findings": fs, "stat": _severity_stat(fs)}]
    buckets: "OrderedDict[str, List[Dict[str, Any]]]" = OrderedDict()
    for f in findings:
        name = f.get("system_name") or "未分组系统"
        buckets.setdefault(name, []).append(dict(f))
    groups = []
    for name, fs in buckets.items():
        fs = _reseq(fs)
        groups.append({"name": name, "findings": fs, "stat": _severity_stat(fs)})
    return groups


def _evidence_text(ev: Any) -> str:
    """evidence 可能是 list[{tool,arguments,result}] 或字符串 → 拼成可读文本。"""
    if isinstance(ev, str):
        return ev
    if isinstance(ev, list):
        parts = []
        for e in ev:
            if isinstance(e, dict):
                res = e.get("result")
                from .report_evidence import _object, _headers
                obj = _object(res)
                if e.get("tool") == "http_request" and obj.get("status_code") is not None:
                    body = obj.get("body", "")
                    if not isinstance(body, str):
                        body = json.dumps(body, ensure_ascii=False, indent=2)
                    label = "留存响应片段" if obj.get("truncated") or obj.get("partial") else "留存响应"
                    parts.append("{} · 状态码 {}\n{}\n\n{}".format(label, obj["status_code"], _headers(obj.get("headers")), body))
                    continue
                res = res if isinstance(res, str) else json.dumps(res, ensure_ascii=False)
                parts.append("[{}] {}".format(e.get("tool", ""), res or ""))
            else:
                parts.append(str(e))
        return "\n".join(parts)
    return ""


def _severity_stat(findings: List[Dict[str, Any]]) -> Dict[str, int]:
    by = {k: 0 for k in _SEV_ORDER}
    for f in findings:
        sv = f.get("severity", "")
        if sv in by:
            by[sv] += 1
    by["total"] = len(findings)
    return by


# ---- 图表（matplotlib Agg headless；CJK 字体防豆腐块）----
_MPL_READY = False


def _ensure_matplotlib():
    """惰性初始化 matplotlib：强制 Agg backend（在 pyplot 前）+ 设中文字体。每进程一次。
    返回 pyplot 模块或 None（未装则降级）。"""
    global _MPL_READY
    try:
        import matplotlib
        matplotlib.use("Agg")   # 必须在 import pyplot 之前（headless，无 GUI 线程问题）
        if not _MPL_READY:
            # 镜像装了 fonts-noto-cjk/fonts-wqy-zenhei；指定字体族避免中文豆腐块。
            matplotlib.rcParams["font.sans-serif"] = [
                "Noto Sans CJK SC", "WenQuanYi Zen Hei", "SimHei", "Microsoft YaHei", "DejaVu Sans"]
            matplotlib.rcParams["axes.unicode_minus"] = False
            _MPL_READY = True
        import matplotlib.pyplot as plt
        return plt
    except Exception as exc:
        logger.debug("matplotlib unavailable: %s", exc)
        return None


_SEV_COLOR = {"critical": "#c0392b", "high": "#e67e22", "medium": "#f1c40f",
              "low": "#3498db", "info": "#95a5a6"}


def _chart_severity(stat: Dict[str, int], out_path: str) -> bool:
    """严重度分布饼图 → PNG。全 0 或无 matplotlib 返 False（生成时跳过该图）。"""
    plt = _ensure_matplotlib()
    if not plt:
        return False
    labels, sizes, colors = [], [], []
    for sv in _SEV_ORDER:
        c = stat.get(sv, 0)
        if c > 0:
            labels.append("{}({})".format(_SEV_CN.get(sv, sv), c))
            sizes.append(c)
            colors.append(_SEV_COLOR.get(sv, "#888"))
    if not sizes:
        return False
    fig = plt.figure(figsize=(5, 4))
    try:
        ax = fig.add_subplot(111)
        ax.pie(sizes, labels=labels, colors=colors, autopct="%1.0f%%", startangle=90)
        ax.set_title("漏洞严重度分布")
        ax.axis("equal")
        fig.savefig(out_path, dpi=120, bbox_inches="tight")
        return True
    except Exception as exc:
        logger.debug("_chart_severity failed: %s", exc)
        return False
    finally:
        plt.close(fig)   # 防内存泄漏 + gthread 状态串扰


def _chart_attack_chain(source: str, source_id: str, out_path: str) -> bool:
    """攻击链竖向流程框图 → PNG（读 intel_attack_chain）。无链/无 matplotlib 返 False。"""
    plt = _ensure_matplotlib()
    if not plt:
        return False
    try:
        q = {"sessions": source_id} if source == "session" else {}
        chains = list(get_repo().collection(Collections.INTEL_ATTACK_CHAIN).find(q).sort("_id", -1).limit(1))
    except Exception:
        chains = []
    if not chains:
        return False
    steps = chains[0].get("steps") or []
    if not steps:
        return False
    steps = steps[:10]   # 上限防超长图
    fig = plt.figure(figsize=(6, max(2, len(steps) * 1.1)))
    try:
        ax = fig.add_subplot(111)
        ax.axis("off")
        n = len(steps)
        for i, s in enumerate(steps):
            y = 1 - (i + 0.5) / n
            label = "{}. {}".format(s.get("seq", i + 1),
                                    (s.get("action") or s.get("vuln_type") or "环节")[:24])
            ax.text(0.5, y, label, ha="center", va="center", fontsize=11,
                    bbox=dict(boxstyle="round,pad=0.4", fc="#eaf2fb", ec="#3498db"))
            if i < n - 1:
                ax.annotate("", xy=(0.5, y - 0.5 / n), xytext=(0.5, y - 0.15 / n),
                            arrowprops=dict(arrowstyle="->", color="#666"))
        ax.set_title("攻击链")
        fig.savefig(out_path, dpi=120, bbox_inches="tight")
        return True
    except Exception as exc:
        logger.debug("_chart_attack_chain failed: %s", exc)
        return False
    finally:
        plt.close(fig)


# ---- 生成主入口 ----

def _abs_template_path(rel_or_abs: str) -> str:
    """schema 里存的是相对 template_dir 的路径，还原成绝对路径。"""
    if os.path.isabs(rel_or_abs):
        return rel_or_abs
    return os.path.join(template_dir(), rel_or_abs)


def _delivery_warnings(data: Dict[str, Any]) -> List[str]:
    gaps = []
    if not data.get("report", {}).get("unit"):
        gaps.append("缺少受影响单位名称")
    for finding in data.get("findings", []):
        if finding.get("verification_status") != "已验证":
            gaps.append("报告包含待验证线索，不能作为已确认漏洞直接提交")
        if not any(s.get("section") == "icp" for s in finding.get("screenshots", [])):
            gaps.append("缺少主体归属证明图片或已保存的备案查询记录")
        for key, label in (("impact", "影响范围"), ("root_cause", "产生原因"), ("remediation", "修复建议"),
                           ("affected_ip", "受影响 IP"), ("discovered_at", "发现时间")):
            if not str(finding.get(key) or "").strip():
                gaps.append("缺少{}，现有证据不足以自动补齐".format(label))
    return list(dict.fromkeys(gaps))


def generate_from_template(template_id: str, source: str, source_id: str,
                           options: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """用模板 + 数据源生成 docx 报告（确定性填充，零/极少 LLM）。
    落 pentest_report（gen_mode=template + docx_path）。返回 {ok, report_id, docx_path} 或 {error}。"""
    options = options or {}
    from . import report_template_ncc as ncc
    if template_id == ncc.TEMPLATE_ID:
        if source != "finding":
            return {"error": "NCC 事件型模板只用于单个漏洞报告"}
        try:
            ncc.seed()
        except Exception as exc:
            return {"error": "NCC 模板准备失败：{}".format(exc)}
    tpl_doc = repo_coll().find_one({"_id": _oid(template_id)})
    if not tpl_doc:
        return {"error": "模板不存在"}
    if tpl_doc.get("status") != "ready":
        return {"error": "模板未就绪（status={}）".format(tpl_doc.get("status"))}
    schema = tpl_doc.get("schema") or {}
    template_path = _abs_template_path(tpl_doc.get("template_path") or "")
    if not os.path.isfile(template_path):
        return {"error": "模板文件缺失：{}".format(template_path)}
    sid = str(source_id or "").strip()
    if not sid:
        return {"error": "source_id 必填"}
    report_key = "tpl:{}:{}".format(sid, template_id)
    previous = get_repo().collection(Collections.PENTEST_REPORT).find_one({"report_key": report_key})
    if previous and previous.get("edited"):
        return {"error": "该报告已有人工修订，请在报告编辑中继续修改；不会用重新生成覆盖修订内容"}

    try:
        from docxtpl import DocxTemplate, InlineImage
        from docx.shared import Mm
    except Exception as exc:
        return {"error": "docx 模板组件未安装（需 rebuild 镜像）：{}".format(exc)}

    findings = _collect_findings(source, sid)
    if not findings:
        return {"error": "没有可生成报告的漏洞记录"}
    try:
        _import_manual_shots(tpl_doc, findings)
    except (ValueError, OSError) as exc:
        return {"error": "补充截图读取失败：{}".format(exc)}
    # 证据截图：把每条漏洞绑定的人工上传截图（finding_shots/<finding_id>/）做成 InlineImage 注入 item.shots，
    # 供模板 {%p for shot in item.shots %}{{ shot }}{%p endfor %} 渲染。放在分组前——systems 的浅拷贝共享同一列表。
    generation_warnings = []
    from .report_evidence import prepare
    for _f in findings:
        generation_warnings.extend(prepare(_f))
    stat = _severity_stat(findings)
    unit, title, system_name = _source_meta(source, sid)
    if source == "finding" and findings:
        unit = unit or findings[0].get("owner_unit", "")
        system_name = findings[0].get("system_name") or system_name
    # 任务级「多一级目录」：按系统分组供内置任务级模板 {%tr for sys in systems%} 渲染。
    # session 源=单组（内置会话级模板不用 systems，忽略即可，多给不碍事）。
    systems = _group_findings_by_system(source, findings, meta_system=system_name)

    ctx: Dict[str, Any] = {"findings": findings, "systems": systems, "stat": stat}
    scalar_vals = {
        "report.title": title, "report.unit": unit, "report.system_name": system_name,
        "report.period": _now()[:10],
        "report.summary": "本次测试共发现漏洞 {} 个。".format(stat.get("total", 0)),
        "stat.total": stat.get("total", 0), "stat.critical": stat.get("critical", 0),
        "stat.high": stat.get("high", 0), "stat.medium": stat.get("medium", 0),
        "stat.low": stat.get("low", 0), "stat.info": stat.get("info", 0),
    }
    for s in schema.get("scalars", []):
        ph = (s.get("placeholder") or "").strip()
        if not ph:
            continue
        ctx[ph] = scalar_vals.get(s.get("semantic") or "", options.get(ph, ""))

    out_root = _tpl_root(template_id)
    os.makedirs(out_root, exist_ok=True)
    chart_paths = {}
    for im in schema.get("images", []):
        ph = (im.get("placeholder") or "").strip()
        if not ph:
            continue
        sem = im.get("semantic") or ""
        png = os.path.join(out_root, "chart_{}_{}.png".format(ph, uuid.uuid4().hex[:8]))
        ok = (_chart_severity(stat, png) if sem == "chart.severity_dist"
              else _chart_attack_chain(source, sid, png) if sem == "chart.attack_chain" else False)
        chart_paths[ph] = png if ok else ""

    out_path = os.path.join(out_root, "out_{}.docx".format(uuid.uuid4().hex))
    report_data = {"source": source, "source_id": sid, "report": {
        "title": title, "unit": unit, "system_name": system_name,
        "period": scalar_vals["report.period"], "summary": scalar_vals["report.summary"],
    }, "findings": findings,
        "variables": {s["placeholder"]: ctx.get(s["placeholder"], "")
                      for s in schema.get("scalars", []) if s.get("placeholder")}}
    if template_id == ncc.TEMPLATE_ID and options.get("ai_assist", True):
        from .report_assist import propose_data
        suggestion = propose_data(report_data, provider_id=str(options.get("provider_id") or ""), only_empty=True)
        if suggestion.get("error"):
            generation_warnings.append(suggestion["error"])
        else:
            by_id = {f["finding_id"]: f for f in findings}
            for change in suggestion.get("changes", []):
                finding = by_id[change["finding_id"]]
                if change["field"] == "step_text":
                    finding["reproduction_steps"][change["step_index"]]["text"] = change["after"]
                else:
                    finding[change["field"]] = change["after"]
            report_data["ai_assistance"] = suggestion
            generation_warnings.extend(suggestion.get("warnings", []))
    if template_id == ncc.TEMPLATE_ID:
        generation_warnings.extend(_delivery_warnings(report_data))
    try:
        # autoescape=True 必需：漏洞 POC/证据常含 < > &（XSS/SQLi payload、HTML 响应体），
        # 不转义会被 docxtpl 当 XML 标签吞掉 → 报告里 POC 内容凭空消失（实测 <script> 整段丢失）。
        from .report_document import render
        render(template_path, report_data, schema, out_path, image_dir(), chart_paths)
    except Exception as exc:
        return {"error": "报告渲染失败：{}".format(exc)}

    max_sev = next((sv for sv in _SEV_ORDER if stat.get(sv, 0) > 0), "")
    now = _now()
    report_key = "tpl:{}:{}".format(sid, template_id)
    _rtype = {"task": "task", "finding": "finding"}.get(source, "session")
    doc = {"report_type": _rtype, "report_key": report_key,
           "gen_mode": "template", "template_id": str(template_id),
           "docx_path": os.path.relpath(out_path, template_dir()).replace("\\", "/"),
           "source_task_id": sid if source == "task" else "",
           "source_session": sid if source == "session" else "",
           "unit": unit, "asset_key": "", "title": title or (tpl_doc.get("name") or "渗透测试报告"),
           "system_name": system_name, "max_severity": max_sev, "vuln_total": stat.get("total", 0),
           "vuln_index": [f["finding_id"] for f in findings if f.get("finding_id")],
           "content": "", "update_date": now}
    doc["report_data"] = report_data
    doc["generation_warnings"] = list(dict.fromkeys(generation_warnings))
    doc["delivery_ready"] = not doc["generation_warnings"]
    rcoll = get_repo().collection(Collections.PENTEST_REPORT)
    existed = rcoll.find_one({"report_key": report_key})
    r = rcoll.update_one({"report_key": report_key},
                         {"$set": doc, "$setOnInsert": {"save_date": now}}, upsert=True)
    rid = existed["_id"] if existed else getattr(r, "upserted_id", None)
    return {"ok": True, "report_id": str(rid) if rid else "",
            "docx_path": doc["docx_path"], "vuln_total": stat.get("total", 0), "updated": bool(existed),
            "generation_warnings": doc["generation_warnings"], "delivery_ready": doc["delivery_ready"]}


def _source_meta(source: str, source_id: str) -> tuple:
    """取报告标题/单位/系统名（从 intel_report 会话报告或 task）。降级空串。"""
    repo = get_repo()
    try:
        if source == "finding":   # 单漏洞级：取该 finding 的 unit/vuln_type 作标题
            f = repo.collection(Collections.INTEL_FINDING).find_one({"_id": _oid(source_id)}) or {}
            vt = f.get("vuln_type") or f.get("name") or "漏洞"
            return f.get("unit", ""), "{} 漏洞报告".format(vt), f.get("system_name", "") or f.get("subdomain", "") or f.get("target", "")
        if source == "session":
            m = repo.collection(Collections.INTEL_REPORT).find_one({"source_session": source_id}) or {}
            return m.get("unit", ""), m.get("title", ""), m.get("system_name", "")
        t = repo.collection(Collections.TASK).find_one({"_id": _oid(source_id)}) if source_id else None
        tname = (t or {}).get("name", "")
        return "", ("{} — 渗透测试报告".format(tname) if tname else ""), ""
    except Exception:
        return "", "", ""


def editor_data(report: Dict[str, Any]) -> Dict[str, Any]:
    """Return a JSON snapshot; legacy template reports can be edited without DB migration."""
    import copy
    if isinstance(report.get("report_data"), dict):
        return copy.deepcopy(report["report_data"])
    source = report.get("report_type") or "session"
    sid = (report.get("source_task_id") if source == "task" else
           (report.get("vuln_index") or [""])[0] if source == "finding" else report.get("source_session"))
    findings = _collect_findings(source, str(sid or "")) if sid else []
    for finding in findings:
        finding["screenshots"] = [{"name": shot["name"], "caption": "", "section": "evidence", "step": 1}
                                  for shot in list_finding_shots(finding["finding_id"])]
    return {"source": source, "source_id": str(sid or ""), "report": {
        "title": report.get("title") or "", "unit": report.get("unit") or "",
        "system_name": report.get("system_name") or "",
        "period": str(report.get("save_date") or "")[:10], "summary": "",
    }, "findings": findings}


def _import_manual_shots(template, findings):
    """Read legacy template-bound images into the same validated evidence store."""
    import hashlib
    from .report_document import normalize_image
    root = os.path.realpath(template_dir())
    for finding in findings:
        fid = finding["finding_id"]
        for rel in (template.get("manual_shots") or {}).get(fid, []):
            path = os.path.realpath(os.path.join(root, rel))
            if os.path.commonpath([root, path]) != root or not os.path.isfile(path):
                raise ValueError("模板绑定的补充截图缺失或路径无效")
            with open(path, "rb") as stream:
                data = normalize_image(stream.read())
            target_dir = _finding_shots_dir(fid)
            os.makedirs(target_dir, exist_ok=True)
            name = "manual_{}.png".format(hashlib.sha256(data).hexdigest()[:24])
            target = os.path.join(target_dir, name)
            if not os.path.isfile(target):
                with open(target, "wb") as stream:
                    stream.write(data)


def update_template_report(report_id: str, report: Dict[str, Any], submitted=None,
                           title=None) -> Dict[str, Any]:
    """Render first, then publish the new path and snapshot together; keep old export on error."""
    from .report_document import render, validate_edit
    current = editor_data(report)
    if not current.get("findings"):
        return {"error": "原报告缺少可编辑的数据快照，且来源漏洞已不存在；原报告未覆盖"}
    try:
        data = validate_edit(current, submitted if submitted is not None else current)
        if title is not None:
            if not isinstance(title, str) or not title.strip():
                return {"error": "报告标题不能为空"}
            data["report"]["title"] = title.strip()
        tid = str(report.get("template_id") or "")
        from . import report_template_ncc as ncc
        if tid == ncc.TEMPLATE_ID:
            ncc.seed()
        template = repo_coll().find_one({"_id": _oid(tid)})
        if not template:
            return {"error": "报告模板不存在，未覆盖原报告"}
        root = _tpl_root(tid)
        os.makedirs(root, exist_ok=True)
        path = os.path.join(root, "edited_{}.docx".format(uuid.uuid4().hex))
        charts = {}
        for im in (template.get("schema") or {}).get("images", []):
            ph = im.get("placeholder", "")
            png = os.path.join(root, "chart_{}.png".format(uuid.uuid4().hex))
            ok = (_chart_severity(_severity_stat(data["findings"]), png)
                  if im.get("semantic") == "chart.severity_dist" else
                  _chart_attack_chain(data["source"], data["source_id"], png)
                  if im.get("semantic") == "chart.attack_chain" else False)
            charts[ph] = png if ok else ""
        render(_abs_template_path(template["template_path"]), data,
               template.get("schema") or {}, path, image_dir(), charts)
        update = {"report_data": data, "title": data["report"].get("title", ""),
                  "unit": data["report"].get("unit", ""), "system_name": data["report"].get("system_name", ""),
                  "docx_path": os.path.relpath(path, template_dir()).replace("\\", "/"),
                  "edited": True, "update_date": _now()}
        stat = _severity_stat(data["findings"])
        update["max_severity"] = next((sv for sv in _SEV_ORDER if stat.get(sv)), "")
        update["generation_warnings"] = _delivery_warnings(data) if tid == ncc.TEMPLATE_ID else []
        update["delivery_ready"] = not update["generation_warnings"]
        result = get_repo().collection(Collections.PENTEST_REPORT).update_one(
            {"_id": _oid(report_id)}, {"$set": update})
        if not getattr(result, "matched_count", 0):
            return {"error": "报告不存在"}
        return {"ok": True, "report_id": report_id, "docx_path": update["docx_path"]}
    except Exception as exc:
        return {"error": "报告修订失败，原报告未覆盖：{}".format(exc)}


# ========== 漏洞证据截图（按 finding_id 全局绑定，跨模板通用）==========
# 存 image_dir()/finding_<finding_id>/<epochms>_<uuid8>.<ext>——落 image 根，复用现有公开
# /api/image/<task_id>/<file> 静态服务（task_id=finding_<fid>），前端 <img> 免 Token 直连预览。
# 不入库（列目录即真相，天然多进程一致）；文件名前缀时间戳保证上传序即渲染序。生成时按 finding 拉图嵌报告。

_IMG_EXT = ("png", "jpg", "jpeg", "gif", "webp")


def _finding_shot_task(finding_id: str) -> str:
    """截图静态服务的 task_id 段（image_dir 下的子目录名）。"""
    return "finding_{}".format(str(finding_id))


def _finding_shots_dir(finding_id: str) -> str:
    from .report_document import safe_component
    return os.path.join(image_dir(), _finding_shot_task(safe_component(finding_id)))


def list_finding_shots(finding_id: str) -> List[Dict[str, str]]:
    """列某漏洞的证据截图（按文件名=上传序）。返回 [{name, url}]，url 走公开 /api/image。目录不存在返空。"""
    fid = (finding_id or "").strip()
    if not fid:
        return []
    d = _finding_shots_dir(fid)
    if not os.path.isdir(d):
        return []
    out: List[Dict[str, str]] = []
    for name in sorted(os.listdir(d)):
        if "." in name and name.lower().rsplit(".", 1)[-1] in _IMG_EXT:
            out.append({"name": name,
                        "url": "/api/image/{}/{}".format(_finding_shot_task(fid), name)})
    return out


def save_finding_shot(finding_id: str, img_bytes: bytes, ext: str = "png") -> Dict[str, Any]:
    """存一张证据截图到 image_dir()/finding_<finding_id>/。返回 {ok, name, url} 或 {error}。"""
    fid = (finding_id or "").strip()
    if not fid:
        return {"error": "finding_id 必填"}
    if not img_bytes:
        return {"error": "空图片"}
    try:
        from .report_document import normalize_image
        img_bytes = normalize_image(img_bytes)
        d = _finding_shots_dir(fid)
    except ValueError as exc:
        return {"error": str(exc)}
    safe_ext = "png"
    os.makedirs(d, exist_ok=True)
    import time as _t
    name = "{}_{}.{}".format(int(_t.time() * 1000), uuid.uuid4().hex[:8], safe_ext)
    try:
        with open(os.path.join(d, name), "wb") as fh:
            fh.write(img_bytes)
    except Exception as exc:
        return {"error": "保存失败：{}".format(exc)}
    return {"ok": True, "name": name, "url": "/api/image/{}/{}".format(_finding_shot_task(fid), name)}


def delete_finding_shot(finding_id: str, name: str) -> Dict[str, Any]:
    """删某漏洞的一张证据截图（防路径遍历：只允许目录内的纯文件名）。"""
    fid = (finding_id or "").strip()
    base = _finding_shots_dir(fid)
    target = os.path.realpath(os.path.join(base, os.path.basename(name or "")))
    if not target.startswith(os.path.realpath(base) + os.sep):
        return {"error": "非法文件名"}
    try:
        if os.path.isfile(target):
            os.remove(target)
            return {"ok": True, "deleted": 1}
        return {"ok": True, "deleted": 0}
    except Exception as exc:
        return {"error": str(exc)}


# ========== 列表 / 详情 / 删除 ==========

def list_templates(page: int = 1, size: int = 20, scope: str = "") -> Dict[str, Any]:
    """模板列表。scope 过滤（v1.21.157-48 item5）：'task'/'session' 只返对应 scope 模板 +
    **无 scope 字段的学习模板**（学习上传的模板不打 scope，两处都可用，不因过滤被藏）。空=全部。"""
    coll = repo_coll()
    try:
        page = max(1, int(page)); size = int(size)
    except (TypeError, ValueError):
        page, size = 1, 20
    q: Dict[str, Any] = {}
    scope = (scope or "").strip()
    if scope in ("task", "session", "finding"):
        # 匹配该 scope 的内置模板 + 无 scope 的学习模板（学习模板通用，各档都能选）
        q["$or"] = [{"scope": scope}, {"scope": {"$exists": False}}, {"scope": ""}, {"scope": None}]
    total = coll.count_documents(q)
    cur = coll.find(q).sort("_id", -1)
    if size and size > 0:
        cur = cur.skip((page - 1) * size).limit(size)
    items = []
    for d in cur:
        d["_id"] = str(d.get("_id", ""))
        d.pop("schema", None)   # 列表不带大 schema，详情才给
        items.append(d)
    return {"items": items, "total": total, "page": page, "size": size}


def get_template(template_id: str) -> Optional[Dict[str, Any]]:
    d = repo_coll().find_one({"_id": _oid(template_id)})
    if not d:
        return None
    d["_id"] = str(d.get("_id", ""))
    return d


def delete_template(template_id: str) -> Dict[str, Any]:
    tid = (template_id or "").strip()
    if not tid:
        return {"error": "template_id 必填"}
    d = repo_coll().find_one({"_id": _oid(tid)})
    if d and d.get("builtin"):
        # 内置默认模板不可删（删了也会在下次启动幂等重播，徒增困惑）——直接拒绝
        return {"error": "内置默认模板不可删除（可复制后修改，或直接用于生成）"}
    try:
        repo_coll().delete_one({"_id": _oid(tid)})
        shutil.rmtree(_tpl_root(tid), ignore_errors=True)   # 连带删磁盘目录
        return {"ok": True, "deleted": 1}
    except Exception as exc:
        return {"error": str(exc)}


# ========== 人工上传补充截图（举证三源之③）==========

def save_manual_shot(template_id: str, finding_id: str, img_bytes: bytes, ext: str = "png") -> Dict[str, Any]:
    """存人工上传的补充截图到 template/<tid>/manual/，绑定到某 finding。返回 {ok, path}。
    绑定关系存 report_template 文档的 manual_shots 映射（finding_id → [相对路径]）。"""
    tid = (template_id or "").strip()
    fid = (finding_id or "").strip()
    if not tid or not fid:
        return {"error": "template_id / finding_id 必填"}
    if not img_bytes:
        return {"error": "空图片"}
    safe_ext = ext.lower().lstrip(".")
    if safe_ext not in ("png", "jpg", "jpeg", "gif", "webp"):
        safe_ext = "png"
    mdir = os.path.join(_tpl_root(tid), "manual")
    os.makedirs(mdir, exist_ok=True)
    fname = "{}.{}".format(uuid.uuid4().hex, safe_ext)
    fpath = os.path.join(mdir, fname)
    try:
        with open(fpath, "wb") as fh:
            fh.write(img_bytes)
    except Exception as exc:
        return {"error": "保存失败：{}".format(exc)}
    rel = os.path.relpath(fpath, template_dir()).replace("\\", "/")
    try:
        repo_coll().update_one({"_id": _oid(tid)},
                               {"$push": {"manual_shots.{}".format(fid): rel}, "$set": {"update_date": _now()}})
    except Exception as exc:
        logger.debug("bind manual_shot degraded: %s", exc)
    return {"ok": True, "path": rel}
